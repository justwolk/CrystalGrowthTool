import os
import openpyxl
from openpyxl.styles import PatternFill
from PySide2.QtGui import QPixmap
from PySide2.QtWidgets import QMessageBox


def save_in_excel(model):
    try:
        if os.path.exists(model.excel_file):
            wb = openpyxl.load_wb(model.excel_file)
        else:
            wb = openpyxl.wb()

        if model.sheet_name in wb.sheetnames:
            sheet = wb[model.sheet_name]
        else:
            sheet = wb.create_sheet(model.sheet_name)
        
        wb.active = sheet
        sheet.delete_rows(1, sheet.max_row)

        all_points = {}
        for i, crystal_obj in enumerate(model.crystal_object):
            for (layer, line_idx), point in crystal_obj.points.items():
                all_points[(i, layer, line_idx)] = point

        image_count = len(model.crystal_object)
        current_row = 0
        nm_value = model.nm_value
        
        current_image_path = os.path.join(model.image_folder, model.photos[0])
        original_pixmap = QPixmap(current_image_path)
        image_width = original_pixmap.width()
        image_height = original_pixmap.height()
        
        all_layers = set()
        all_line_indices = set()
        for (i, layer, line_idx) in all_points.keys():
            all_layers.add(layer)
            all_line_indices.add(line_idx)
        
        for layer in sorted(all_layers):
            layer_points_exist = any((i, layer, line_idx) in all_points for i in range(image_count) for line_idx in all_line_indices)
            if layer_points_exist:
                current_row += 2
                sheet.cell(row=current_row, column=1, value=f"Ступень {layer+1}")
                
                grey_fill = PatternFill(start_color='757575', end_color='757575', fill_type='solid')
                max_excel_column = image_count * 2 + 3
                for col in range(1, max_excel_column + 1):
                    sheet.cell(row=current_row, column=col).fill = grey_fill

                for line_idx in sorted(all_line_indices):
                    line_points_exist = any((i, layer, line_idx) in all_points for i in range(image_count))
                    if line_points_exist:
                        current_row += 1
                        if model.save_half:
                            sheet.cell(row=current_row, column=1, value=float(line_idx+1)/2)
                        else:
                            sheet.cell(row=current_row, column=1, value=float(line_idx+1))

                        for i in range(image_count):
                            if (i, layer, line_idx) in all_points:
                                point = all_points[(i, layer, line_idx)]
                                
                                x_nm = (point.x() / image_width) * nm_value
                                y_nm = ((image_height - point.y()) / image_height) * nm_value
                                
                                x_nm = max(0, min(nm_value, x_nm))
                                y_nm = max(0, min(nm_value, y_nm))
                                
                                accuracy = getattr(model, 'excel_save_accuracy', 2)

                                sheet.cell(row=current_row, column=i * 2 + 2, value=round(x_nm, accuracy))
                                sheet.cell(row=current_row, column=i * 2 + 3, value=round(y_nm, accuracy))

                                sheet.cell(row=1, column=i * 2 + 3, value=str(int(i)+1))

        wb.save(model.excel_file)

    except PermissionError:
        msg_box = QMessageBox()
        msg_box.setIcon(QMessageBox.Warning)
        msg_box.setWindowTitle("Ошибка доступа")
        msg_box.setText("Файл эксель открыт, сначала закройте его и повторите попытку.")
        msg_box.exec_()

    except Exception as e:
        print(f"Ошибка: {e}")
